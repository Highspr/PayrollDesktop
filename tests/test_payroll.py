import importlib.util
import sqlite3
import time
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from docx import Document

from payroll_app import core
from payroll_app.api import create_app
from payroll_app.service import Store


@pytest.fixture
def workspace(tmp_path):
    store = Store(tmp_path / 'data')
    client = TestClient(create_app(store, 'test-secret'), headers={'X-Payroll-Token': 'test-secret'})
    employee = {'name':'Test Teacher', 'bps':17, 'basic':30000, 'special':2000, 'account_no':'00123456789', 'bank':'ABL'}
    eid = client.post('/api/employees', json=employee).json()['id']
    client.post('/api/scales', json={'bps':17,'hr':5000,'conv':4000,'medl':1000,'adhoc':3000,'cpf':3000})
    return store, client, eid


def test_formula():
    row = core.calculate(30000, {'hr':5000,'conv':4000,'medl':1000,'adhoc':3000,'cpf':3000},2000,True,core.DEFAULTS,30,28,2000,1000,500)
    assert row['gross']==47000
    assert row['lwop_pay']==3000
    assert row['net']==39500
    assert row['cpf_college']==3000
    contract=core.calculate(30000,None,0,False,core.DEFAULTS,30,30)
    assert contract['net']==30000


def test_month_edits_lock_and_validation(workspace):
    store, client, eid = workspace
    row=client.get('/api/payroll/2026-08').json()['rows'][0]
    row.update(worked=28,arrear=2000,cp_loan=1000,other=500,des='Lecturer')
    assert client.put(f'/api/payroll/2026-08/{eid}',json=row).status_code==200
    august=client.get('/api/payroll/2026-08').json()
    assert august['rows'][0]['net']==39500
    assert august['rows'][0]['des']=='Lecturer'
    assert client.get('/api/payroll/2026-09').json()['rows'][0]['worked']==30
    assert client.post('/api/payroll/2026-08/lock').status_code==200
    assert client.put(f'/api/payroll/2026-08/{eid}',json=row).status_code==409
    client.post('/api/scales',json={'bps':17,'hr':9999})
    assert client.get('/api/payroll/2026-08').json()=={**august,'locked':True}
    assert client.post('/api/payroll/2026-08/unlock').status_code==200
    row['worked']=32
    assert client.put(f'/api/payroll/2026-08/{eid}',json=row).status_code==422
    row['worked']=29;row['wd']=28
    assert client.put(f'/api/payroll/2026-08/{eid}',json=row).status_code==422


def test_auth_and_origin(workspace):
    store, client, eid=workspace
    assert client.get('/api/employees',headers={'X-Payroll-Token':'wrong'}).status_code==401
    assert client.get('/api/employees',headers={'Origin':'https://evil.test'}).status_code==403
    assert client.get('/api/employees',headers={'Host':'evil.test'}).status_code==403
    assert client.get('/api/payroll/2026-99').status_code==400


def test_legacy_employee_null_fields_can_be_edited(workspace):
    store,client,eid=workspace
    with store.db() as db:
        db.execute('UPDATE employees SET email=NULL, section=NULL, designation=NULL WHERE id=?',(eid,))
        db.commit()
    row=client.get('/api/employees').json()[0]
    row['name']='Updated name'
    assert client.put(f'/api/employees/{eid}',json=row).status_code==200


def test_exports_backup_restore(workspace,tmp_path):
    store,client,eid=workspace
    for kind,extension in [('pdf','pdf'),('salary','xlsx'),('bank','xlsx'),('cpf','xlsx'),('word','docx')]:
        out=tmp_path/f'{kind}.{extension}'
        response=client.post('/api/export',json={'month':'2026-08','kind':kind,'ticket':store.ticket(out,'export')})
        assert response.status_code==200,response.text
        assert out.stat().st_size>500
        if extension=='xlsx':
            wb=load_workbook(out)
            assert wb.active.max_row>3
        elif extension=='docx':
            assert Document(out).tables[0].cell(1,0).text=='Test Teacher'
        else:
            assert out.read_bytes().startswith(b'%PDF')
    backup=tmp_path/'backup.db'
    assert client.post('/api/backup',json={'ticket':store.ticket(backup,'backup')}).status_code==200
    client.post('/api/employees',json={'name':'Second Teacher'})
    assert len(client.get('/api/employees').json())==2
    assert client.post('/api/restore',json={'ticket':store.ticket(backup,'restore')}).status_code==200
    assert len(client.get('/api/employees').json())==1
    assert len(list((store.directory/'backups').glob('*.db')))>=1
    wrong=tmp_path/'wrong.db'
    with sqlite3.connect(wrong) as db:
        db.execute('create table junk (id integer)')
    assert client.post('/api/restore',json={'ticket':store.ticket(wrong,'restore')}).status_code==400
    assert len(client.get('/api/employees').json())==1


def test_migration_matches_legacy(tmp_path):
    original=Path(__file__).resolve().parents[1]/'payroll.db'
    if not original.exists():
        pytest.skip('No legacy fixture present')
    before=original.read_bytes()
    store=Store(tmp_path/'migrated',original)
    spec=importlib.util.spec_from_file_location('legacy',str(original.with_suffix('.py')))
    legacy=importlib.util.module_from_spec(spec)
    spec.loader.exec_module(legacy)
    with sqlite3.connect(original.as_uri()+'?mode=ro',uri=True) as db:
        db.row_factory=sqlite3.Row
        expected=legacy.month_rows(db,'2026-08')
    with store.db() as db:
        assert core.month_rows(db,'2026-08')==expected
    assert original.read_bytes()==before


def test_400_staff(workspace):
    store,client,eid=workspace
    with store.db() as db:
        db.executemany('INSERT INTO employees(name,basic,bps) VALUES (?,?,?)',[(f'Teacher {i}',30000,17) for i in range(399)])
        db.commit()
    started=time.perf_counter()
    result=client.get('/api/payroll/2026-08').json()
    assert len(result['rows'])==400
    assert time.perf_counter()-started<2


def test_payslip_long_names_and_400_pages(workspace,tmp_path):
    import pymupdf
    store,client,eid=workspace
    with store.db() as db:
        rows,_=core.month_rows(db,'2026-08')
    rows[0].update(name='Professor Muhammad Abdullah Khan ' * 5,
                   des='Senior Lecturer, Department of Computer Science and Information Technology',
                   acct='PK00ABCD000012345678901234',section='Senior College - Computing and Science')
    settings=dict(core.DEFAULTS, college_name='College of Computing and Applied Sciences',
                  slip_header='Main Campus, College Road, Islamabad - Accounts and Administration Department',
                  signatory_name='Accounts Officer',signatory_title='Finance and Administration')
    path=tmp_path/'batch.pdf'
    started=time.perf_counter()
    core.make_payslips(str(path),rows*400,settings,'August 2026 - DRAFT')
    elapsed=time.perf_counter()-started
    with pymupdf.open(path) as pdf:
        assert len(pdf)==400
        assert 'PK00ABCD000012345678901234' in pdf[0].get_text().replace('\n','')
        assert '42,000' in pdf[0].get_text()
        assert all(0<=b[0]<b[2]<=pdf[0].rect.width and 0<=b[1]<b[3]<=pdf[0].rect.height for b in pdf[0].get_text('blocks'))
    assert elapsed<60


def test_import_is_staged_and_locked_month_protected(workspace,tmp_path):
    store,client,eid=workspace
    source=Path(__file__).resolve().parents[1]/'salary data.xlsx'
    if not source.exists():
        pytest.skip('No college workbook fixture')
    result=client.post('/api/import',json={'ticket':store.ticket(source,'import')})
    assert result.status_code==200,result.text
    month=f"{result.json()['year']}-{result.json()['month']:02}"
    # Use a direct lock fixture: existing college data may require validation first.
    with store.db() as db:
        rows,_=core.month_rows(db,month)
        core.lock_month(db,month,rows)
    before=store.path.read_bytes()
    assert client.post('/api/import',json={'ticket':store.ticket(source,'import')}).status_code==409
    assert store.path.read_bytes()==before
